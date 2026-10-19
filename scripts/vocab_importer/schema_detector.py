import warnings

def detect_schema(filepath: str):
    warnings.simplefilter("ignore")
    try:
        # Keep identity/metadata imports usable in environments that do not
        # install the optional spreadsheet parser. Workbook parsing still
        # reports the missing dependency as a structured parse failure.
        from openpyxl import load_workbook

        wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(max_row=20, values_only=True))
        if not rows:
            return "EMPTY", None
            
        header = [str(c).strip().lower() if c else "" for c in rows[0]]
        
        # SCHEMA_B: 单词 英音 美音 释义
        if "单词" in header and "释义" in header:
            if "英音" in header or "美音" in header:
                return "SCHEMA_B", header
        
        # SCHEMA_A: word phonetic definition
        if "word" in header and "definition" in header:
            return "SCHEMA_A", header
            
        # 检查是否是无表头 SCHEMA_HEADERLESS
        # 第一列全是英文，第二列全是包含中文的
        eng_col0 = 0
        cn_col1 = 0
        valid_rows = 0
        for r in rows:
            if len(r) >= 2 and r[0] and r[1]:
                val0 = str(r[0]).strip()
                val1 = str(r[1]).strip()
                if val0 and not val0.isdigit():
                    valid_rows += 1
                    # 简单判断是否全英文字符(允许空格和')
                    if all(ord(c) < 128 for c in val0):
                        eng_col0 += 1
                    if any(ord(c) > 255 for c in val1):
                        cn_col1 += 1
        
        if valid_rows > 0 and (eng_col0 / valid_rows) > 0.8 and (cn_col1 / valid_rows) > 0.8:
            return "SCHEMA_HEADERLESS", None
            
        return "UNKNOWN_SCHEMA", header
    except Exception as e:
        return "FATAL_PARSE_ERROR", str(e)
        
